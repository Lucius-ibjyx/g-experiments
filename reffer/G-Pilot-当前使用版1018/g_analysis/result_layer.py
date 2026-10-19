from __future__ import annotations

from pathlib import Path
import math
import re
from typing import Iterable

import numpy as np
import pandas as pd


def _bool_series(s: pd.Series) -> pd.Series:
    if s.dtype == bool:
        return s.fillna(False)
    return s.astype(str).str.lower().map({"true": True, "false": False, "1": True, "0": False}).fillna(False)


def _safe_name(text: str) -> str:
    text = str(text)
    text = re.sub(r"[^0-9A-Za-z._-]+", "_", text)
    return text.strip("_") or "unnamed"



def _eccentricity_series(df: pd.DataFrame) -> pd.Series:
    """Return exact eccentricity when available; otherwise reconstruct from gaze geometry."""
    if "eccentricity_deg" in df.columns:
        return pd.to_numeric(df["eccentricity_deg"], errors="coerce")

    xyz = {"GazePositionX", "GazePositionY", "GazePositionZ"}
    if xyz.issubset(df.columns):
        x = pd.to_numeric(df["GazePositionX"], errors="coerce")
        y = pd.to_numeric(df["GazePositionY"], errors="coerce")
        z = pd.to_numeric(df["GazePositionZ"], errors="coerce")
        return pd.Series(
            np.degrees(np.arctan2(np.sqrt(x * x + y * y), z)),
            index=df.index,
            dtype=float,
        )

    h = pd.to_numeric(df.get("horizontal_visual_angle_deg"), errors="coerce")
    v = pd.to_numeric(df.get("vertical_visual_angle_deg"), errors="coerce")
    hr = np.radians(h)
    vr = np.radians(v)
    return pd.Series(
        np.degrees(np.arctan(np.sqrt(np.tan(hr) ** 2 + np.tan(vr) ** 2))),
        index=df.index,
        dtype=float,
    )

def _sample_period_s(fused: pd.DataFrame) -> float:
    t = pd.to_numeric(fused.get("video_time_s"), errors="coerce").dropna().sort_values()
    d = t.diff().dropna()
    d = d[(d > 0) & (d < 0.1)]
    if len(d):
        return float(d.median())
    return 1.0 / 70.0


def _deduplicate_semantic_hits(long_hits: pd.DataFrame, fused: pd.DataFrame) -> pd.DataFrame:
    """One hit row per sample_id + semantic_key.

    Physical CVAT tracks can be separate while analysis semantics are merged.
    For example, protruding and centered A520 doorplates both map to target__A520.
    If the same semantic object has multiple physical-track hits at one sample,
    keep exact over tolerance, then the smallest boundary distance.
    """
    if long_hits.empty:
        cols = list(long_hits.columns) + ["stage_id", "stage_type", "stage_occurrence"]
        return pd.DataFrame(columns=list(dict.fromkeys(cols)))

    h = long_hits.copy()
    stage_cols = [c for c in ["sample_id", "stage_id", "stage_type", "stage_occurrence"] if c in fused.columns]
    h = h.merge(fused[stage_cols].drop_duplicates("sample_id"), on="sample_id", how="inner")
    h = h[h["stage_id"].notna()].copy()
    h["hit_rank"] = h["hit_type"].map({"exact": 0, "tolerance": 1}).fillna(2)
    h["boundary_distance_deg"] = pd.to_numeric(h["boundary_distance_deg"], errors="coerce")
    h = h.sort_values(["sample_id", "semantic_key", "hit_rank", "boundary_distance_deg", "track_id"])
    h = h.drop_duplicates(["sample_id", "semantic_key"], keep="first")
    return h.drop(columns=["hit_rank"])


def _build_semantic_episodes(hits: pd.DataFrame, gap_s: float, sample_period_s: float) -> pd.DataFrame:
    """Build AOI access episodes for each semantic AOI, preserving exact/tolerance evidence.

    This is an AOI-access episode, NOT a fixation. Consecutive hits to the same semantic AOI
    are merged when the time gap is <= gap_s and the Stage is unchanged.
    """
    cols = [
        "participant_id", "task_id", "stage_id", "stage_type", "stage_occurrence",
        "semantic_key", "standard_class", "target_status", "room_id",
        "visit_index_within_semantic", "start_s", "end_s", "duration_s",
        "sample_count", "exact_sample_count", "tolerance_sample_count", "visit_hit_type",
        "min_boundary_distance_deg", "track_ids"
    ]
    if hits.empty:
        return pd.DataFrame(columns=cols)

    rows = []
    h = hits.sort_values(["semantic_key", "video_time_s", "sample_id"]).copy()
    for semantic, g in h.groupby("semantic_key", dropna=False, sort=False):
        g = g.sort_values("video_time_s")
        current = []
        last_t = None
        last_stage = None
        visit_idx = 0

        def flush(buf):
            nonlocal visit_idx
            if not buf:
                return
            visit_idx += 1
            d = pd.DataFrame(buf)
            start = float(d["video_time_s"].min())
            end = float(d["video_time_s"].max())
            exact_n = int((d["hit_type"] == "exact").sum())
            tol_n = int((d["hit_type"] == "tolerance").sum())
            hit_type = "exact" if exact_n > 0 else "tolerance"
            dist = pd.to_numeric(d["boundary_distance_deg"], errors="coerce")
            min_dist = float(dist.min()) if dist.notna().any() else np.nan
            row0 = d.iloc[0]
            rows.append({
                "participant_id": row0.get("participant_id", pd.NA),
                "task_id": row0.get("task_id", pd.NA),
                "stage_id": row0.get("stage_id", pd.NA),
                "stage_type": row0.get("stage_type", pd.NA),
                "stage_occurrence": row0.get("stage_occurrence", pd.NA),
                "semantic_key": semantic,
                "standard_class": row0.get("standard_class", pd.NA),
                "target_status": row0.get("target_status", pd.NA),
                "room_id": row0.get("room_id", pd.NA),
                "visit_index_within_semantic": visit_idx,
                "start_s": start,
                "end_s": end,
                "duration_s": max(sample_period_s, end - start + sample_period_s),
                "sample_count": int(len(d)),
                "exact_sample_count": exact_n,
                "tolerance_sample_count": tol_n,
                "visit_hit_type": hit_type,
                "min_boundary_distance_deg": min_dist,
                "track_ids": "|".join(map(str, sorted(set(pd.to_numeric(d["track_id"], errors="coerce").dropna().astype(int))))),
            })

        for _, r in g.iterrows():
            t = float(r["video_time_s"])
            stage = (r.get("stage_id"), r.get("stage_type"), r.get("stage_occurrence"))
            if current and (last_t is None or t - last_t > gap_s or stage != last_stage):
                flush(current)
                current = []
            current.append(r.to_dict())
            last_t = t
            last_stage = stage
        flush(current)

    return pd.DataFrame(rows, columns=cols)


def _build_primary_episodes(fused: pd.DataFrame, gap_s: float, sample_period_s: float) -> pd.DataFrame:
    """Build a non-overlapping primary-AOI sequence for transition analysis."""
    cols = [
        "sequence_index", "participant_id", "task_id", "stage_id", "stage_type", "stage_occurrence",
        "semantic_key", "standard_class", "start_s", "end_s", "duration_s", "sample_count",
        "visit_hit_type", "min_boundary_distance_deg"
    ]
    g = fused.copy()
    valid = _bool_series(g["gaze_valid"]) if "gaze_valid" in g.columns else pd.Series(False, index=g.index)
    g = g[valid & g["aoi_primary_semantic_key"].notna()].sort_values("video_time_s")
    if g.empty:
        return pd.DataFrame(columns=cols)

    rows = []
    current = []
    last_t = None
    last_sem = None
    last_stage = None
    seq = 0

    def flush(buf):
        nonlocal seq
        if not buf:
            return
        seq += 1
        d = pd.DataFrame(buf)
        start = float(d["video_time_s"].min())
        end = float(d["video_time_s"].max())
        types = set(d["aoi_primary_hit_type"].dropna().astype(str))
        hit_type = "exact" if "exact" in types else "tolerance"
        dist = pd.to_numeric(d["aoi_primary_boundary_distance_deg"], errors="coerce")
        row0 = d.iloc[0]
        rows.append({
            "sequence_index": seq,
            "participant_id": row0.get("participant_id", pd.NA),
            "task_id": row0.get("task_id", pd.NA),
            "stage_id": row0.get("stage_id", pd.NA),
            "stage_type": row0.get("stage_type", pd.NA),
            "stage_occurrence": row0.get("stage_occurrence", pd.NA),
            "semantic_key": row0.get("aoi_primary_semantic_key", pd.NA),
            "standard_class": row0.get("aoi_primary_class", pd.NA),
            "start_s": start,
            "end_s": end,
            "duration_s": max(sample_period_s, end - start + sample_period_s),
            "sample_count": int(len(d)),
            "visit_hit_type": hit_type,
            "min_boundary_distance_deg": float(dist.min()) if dist.notna().any() else np.nan,
        })

    for _, r in g.iterrows():
        t = float(r["video_time_s"])
        sem = r["aoi_primary_semantic_key"]
        stage = (r.get("stage_id"), r.get("stage_type"), r.get("stage_occurrence"))
        if current and (sem != last_sem or stage != last_stage or (last_t is not None and t - last_t > gap_s)):
            flush(current)
            current = []
        current.append(r.to_dict())
        last_t, last_sem, last_stage = t, sem, stage
    flush(current)
    return pd.DataFrame(rows, columns=cols)


def _qc_summary(fused: pd.DataFrame, stages: pd.DataFrame, stage_qc: pd.DataFrame) -> pd.DataFrame:
    rows = []
    valid = _bool_series(fused["gaze_valid"]) if "gaze_valid" in fused.columns else pd.Series(False, index=fused.index)
    rows.append({
        "scope": "whole_recording",
        "stage_id": "",
        "stage_type": "",
        "duration_s": float(pd.to_numeric(fused["video_time_s"], errors="coerce").max() - pd.to_numeric(fused["video_time_s"], errors="coerce").min()) if len(fused) else np.nan,
        "samples": int(len(fused)),
        "valid_samples": int(valid.sum()),
        "valid_gaze_ratio": float(valid.mean()) if len(fused) else np.nan,
        "positive_duration": "",
        "within_video": "",
        "gap_from_previous_s": "",
    })
    qmap = stage_qc.set_index("stage_id").to_dict("index") if not stage_qc.empty else {}
    for _, st in stages.iterrows():
        g = fused[fused["stage_id"] == st["stage_id"]]
        v = _bool_series(g["gaze_valid"]) if len(g) else pd.Series(dtype=bool)
        q = qmap.get(st["stage_id"], {})
        rows.append({
            "scope": "stage",
            "stage_id": st["stage_id"],
            "stage_type": st["stage_type"],
            "duration_s": float(st["duration_s"]),
            "samples": int(len(g)),
            "valid_samples": int(v.sum()) if len(g) else 0,
            "valid_gaze_ratio": float(v.mean()) if len(g) else np.nan,
            "positive_duration": q.get("positive_duration", ""),
            "within_video": q.get("within_video", ""),
            "gap_from_previous_s": q.get("gap_from_previous_s", ""),
        })
    return pd.DataFrame(rows)


def _stage_summary(fused: pd.DataFrame, stages: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for _, st in stages.iterrows():
        g = fused[fused["stage_id"] == st["stage_id"]].copy()
        vmask = _bool_series(g["gaze_valid"]) if len(g) else pd.Series(dtype=bool)
        gv = g[vmask].copy() if len(g) else g
        spatial=g[_bool_series(g["spatial_valid"])] if "spatial_valid" in g else gv
        h = pd.to_numeric(spatial.get("horizontal_visual_angle_deg"), errors="coerce").dropna()
        v = pd.to_numeric(spatial.get("vertical_visual_angle_deg"), errors="coerce").dropna()
        ecc = _eccentricity_series(spatial).dropna()
        primary_hits = gv["aoi_primary_semantic_key"].notna().sum() if "aoi_primary_semantic_key" in gv.columns else 0
        has_aoi=bool(g["aoi_available"].any()) if "aoi_available" in g else True
        rows.append({
            "stage_id": st["stage_id"],
            "stage_name_cn": st.get("stage_name_cn", ""),
            "stage_type": st["stage_type"],
            "occurrence": st.get("occurrence", 1),
            "search_target": st.get("search_target", ""),
            "search_result": st.get("search_result", ""),
            "start_s": float(st["start_s"]),
            "end_s": float(st["end_s"]),
            "duration_s": float(st["duration_s"]),
            "gaze_samples": int(len(g)),
            "valid_gaze_samples": int(len(gv)),
            "valid_spatial_samples":int(len(spatial)),
            "valid_gaze_ratio": float(len(gv) / len(g)) if len(g) else np.nan,
            "aoi_available":has_aoi,
            "primary_aoi_hit_samples": int(primary_hits) if has_aoi else np.nan,
            "primary_aoi_hit_ratio_of_valid": float(primary_hits / len(gv)) if len(gv) and has_aoi else np.nan,
            "h_median_deg": float(h.median()) if len(h) else np.nan,
            "h_p05_deg": float(h.quantile(0.05)) if len(h) else np.nan,
            "h_p95_deg": float(h.quantile(0.95)) if len(h) else np.nan,
            "h_span90_deg": float(h.quantile(0.95) - h.quantile(0.05)) if len(h) else np.nan,
            "v_median_deg": float(v.median()) if len(v) else np.nan,
            "v_p05_deg": float(v.quantile(0.05)) if len(v) else np.nan,
            "v_p95_deg": float(v.quantile(0.95)) if len(v) else np.nan,
            "v_span90_deg": float(v.quantile(0.95) - v.quantile(0.05)) if len(v) else np.nan,
            "eccentricity_median_deg": float(ecc.median()) if len(ecc) else np.nan,
            "eccentricity_p90_deg": float(ecc.quantile(0.90)) if len(ecc) else np.nan,
            "eccentricity_p95_deg": float(ecc.quantile(0.95)) if len(ecc) else np.nan,
        })
    return pd.DataFrame(rows)


def _semantic_inventory(tracks: pd.DataFrame) -> pd.DataFrame:
    if tracks.empty:
        return pd.DataFrame(columns=["semantic_key", "standard_class", "target_status", "room_id", "physical_track_count", "physical_track_ids", "raw_labels"])
    t = tracks[tracks["standard_class"] != "unmapped"].copy()
    rows = []
    for sem, g in t.groupby("semantic_key", dropna=False, sort=False):
        rows.append({
            "semantic_key": sem,
            "standard_class": g["standard_class"].dropna().iloc[0] if g["standard_class"].notna().any() else pd.NA,
            "target_status": g["target_status"].dropna().iloc[0] if g["target_status"].notna().any() else pd.NA,
            "room_id": g["room_id"].dropna().iloc[0] if g["room_id"].notna().any() else pd.NA,
            "physical_track_count": int(g["track_id"].nunique()),
            "physical_track_ids": "|".join(map(str, sorted(pd.to_numeric(g["track_id"], errors="coerce").dropna().astype(int).unique()))),
            "raw_labels": "|".join(sorted(g["raw_label"].dropna().astype(str).unique())),
        })
    return pd.DataFrame(rows)


def _aoi_summary(episodes: pd.DataFrame, tracks: pd.DataFrame) -> pd.DataFrame:
    cols = [
        "semantic_key", "standard_class", "target_status", "room_id", "physical_track_count", "physical_track_ids", "raw_labels", "visited",
        "first_access_s", "last_access_s", "visit_count", "revisit_count", "total_visit_duration_s",
        "exact_visit_count", "tolerance_only_visit_count", "total_hit_samples", "exact_hit_samples",
        "tolerance_hit_samples", "min_boundary_distance_deg"
    ]
    inv = _semantic_inventory(tracks)
    rows = []
    visited_map = {sem: g for sem, g in episodes.groupby("semantic_key", dropna=False, sort=False)} if not episodes.empty else {}
    for _, item in inv.iterrows():
        sem = item["semantic_key"]
        g = visited_map.get(sem)
        base = item.to_dict()
        if g is None or g.empty:
            rows.append({
                **base, "visited": False, "first_access_s": np.nan, "last_access_s": np.nan,
                "visit_count": 0, "revisit_count": 0, "total_visit_duration_s": 0.0,
                "exact_visit_count": 0, "tolerance_only_visit_count": 0, "total_hit_samples": 0,
                "exact_hit_samples": 0, "tolerance_hit_samples": 0, "min_boundary_distance_deg": np.nan,
            })
        else:
            rows.append({
                **base, "visited": True,
                "first_access_s": float(g["start_s"].min()), "last_access_s": float(g["end_s"].max()),
                "visit_count": int(len(g)), "revisit_count": int(max(0, len(g) - 1)),
                "total_visit_duration_s": float(g["duration_s"].sum()),
                "exact_visit_count": int((g["visit_hit_type"] == "exact").sum()),
                "tolerance_only_visit_count": int((g["visit_hit_type"] == "tolerance").sum()),
                "total_hit_samples": int(g["sample_count"].sum()),
                "exact_hit_samples": int(g["exact_sample_count"].sum()),
                "tolerance_hit_samples": int(g["tolerance_sample_count"].sum()),
                "min_boundary_distance_deg": float(pd.to_numeric(g["min_boundary_distance_deg"], errors="coerce").min()),
            })
    return pd.DataFrame(rows, columns=cols).sort_values(["visited", "first_access_s"], ascending=[False, True], na_position="last")


def _aoi_by_stage(episodes: pd.DataFrame) -> pd.DataFrame:
    if episodes.empty:
        return pd.DataFrame(columns=["stage_id", "stage_type", "semantic_key", "standard_class", "visit_count", "total_visit_duration_s", "exact_visit_count", "tolerance_only_visit_count"])
    rows = []
    keys = ["stage_id", "stage_type", "semantic_key", "standard_class"]
    for key, g in episodes.groupby(keys, dropna=False, sort=False):
        rows.append({
            **dict(zip(keys, key)),
            "visit_count": int(len(g)),
            "total_visit_duration_s": float(g["duration_s"].sum()),
            "exact_visit_count": int((g["visit_hit_type"] == "exact").sum()),
            "tolerance_only_visit_count": int((g["visit_hit_type"] == "tolerance").sum()),
            "first_access_s": float(g["start_s"].min()),
            "last_access_s": float(g["end_s"].max()),
        })
    return pd.DataFrame(rows).sort_values(["stage_id", "first_access_s"])


def _transitions(primary_episodes: pd.DataFrame) -> pd.DataFrame:
    cols = ["from_semantic_key", "to_semantic_key", "transition_count"]
    if primary_episodes.empty or len(primary_episodes) < 2:
        return pd.DataFrame(columns=cols)
    p = primary_episodes.sort_values("sequence_index").copy()
    pairs = []
    prev = None
    prev_stage = None
    for _, r in p.iterrows():
        cur = r["semantic_key"]
        stage=(r.get("stage_id"),r.get("stage_type"))
        if prev is not None and stage == prev_stage and cur != prev:
            pairs.append((prev, cur))
        prev = cur
        prev_stage=stage
    if not pairs:
        return pd.DataFrame(columns=cols)
    d = pd.DataFrame(pairs, columns=["from_semantic_key", "to_semantic_key"])
    return d.value_counts().reset_index(name="transition_count").sort_values("transition_count", ascending=False)


def _doorplate_summary(episodes: pd.DataFrame, tracks: pd.DataFrame) -> pd.DataFrame:
    inv = _semantic_inventory(tracks)
    inv = inv[inv["standard_class"].isin(["target_doorplate", "non_target_doorplate"])].copy()
    cols = [
        "target_status", "room_id", "semantic_key", "physical_track_count", "physical_track_ids", "raw_labels",
        "visited", "first_access_s", "last_access_s", "visit_count", "revisit_count", "total_visit_duration_s",
        "exact_visit_count", "tolerance_only_visit_count", "total_hit_samples", "min_boundary_distance_deg"
    ]
    rows = []
    visited_map = {sem: g for sem, g in episodes.groupby("semantic_key", dropna=False, sort=False)} if not episodes.empty else {}
    for _, item in inv.iterrows():
        sem = item["semantic_key"]
        g = visited_map.get(sem)
        base = {
            "target_status": item["target_status"], "room_id": item["room_id"], "semantic_key": sem,
            "physical_track_count": item["physical_track_count"], "physical_track_ids": item["physical_track_ids"], "raw_labels": item["raw_labels"],
        }
        if g is None or g.empty:
            rows.append({
                **base, "visited": False, "first_access_s": np.nan, "last_access_s": np.nan,
                "visit_count": 0, "revisit_count": 0, "total_visit_duration_s": 0.0,
                "exact_visit_count": 0, "tolerance_only_visit_count": 0, "total_hit_samples": 0,
                "min_boundary_distance_deg": np.nan,
            })
        else:
            rows.append({
                **base, "visited": True, "first_access_s": float(g["start_s"].min()), "last_access_s": float(g["end_s"].max()),
                "visit_count": int(len(g)), "revisit_count": int(max(0, len(g) - 1)),
                "total_visit_duration_s": float(g["duration_s"].sum()),
                "exact_visit_count": int((g["visit_hit_type"] == "exact").sum()),
                "tolerance_only_visit_count": int((g["visit_hit_type"] == "tolerance").sum()),
                "total_hit_samples": int(g["sample_count"].sum()),
                "min_boundary_distance_deg": float(pd.to_numeric(g["min_boundary_distance_deg"], errors="coerce").min()),
            })
    return pd.DataFrame(rows, columns=cols).sort_values(["target_status", "room_id"], na_position="last")


def _style_excel(path: Path):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col_idx, col in enumerate(ws.iter_cols(1, ws.max_column), start=1):
            max_len = 0
            for cell in col[: min(len(col), 200)]:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 30)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)
    wb.save(path)


def _make_figures(fused: pd.DataFrame, stages: pd.DataFrame, primary_episodes: pd.DataFrame,
                  outdir: Path, scene_hfov_deg: float, scene_vfov_deg: float):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    available={f.name for f in font_manager.fontManager.ttflist}
    for font in ['PingFang SC','Heiti SC','Microsoft YaHei','SimHei','Noto Sans CJK SC','Arial Unicode MS']:
        if font in available:
            plt.rcParams['font.sans-serif']=[font,'DejaVu Sans'];break
    plt.rcParams['axes.unicode_minus']=False

    figdir = outdir / "figures"
    figdir.mkdir(exist_ok=True)
    valid_all=fused[_bool_series(fused["spatial_valid"] if "spatial_valid" in fused else fused["gaze_valid"])]
    hh=pd.to_numeric(valid_all["horizontal_visual_angle_deg"],errors="coerce").dropna()
    vv=pd.to_numeric(valid_all["vertical_visual_angle_deg"],errors="coerce").dropna()
    hl=max(5,float(hh.abs().max())+2) if len(hh) else 5
    vl=max(5,float(vv.abs().max())+2) if len(vv) else 5
    full_extent=[-hl,hl,-vl,vl]

    # One H/V density figure per Stage. Fixed axes make within-task comparison direct.
    for _, st in stages.iterrows():
        g = fused[(fused["stage_id"] == st["stage_id"]) & _bool_series(fused["spatial_valid"] if "spatial_valid" in fused else fused["gaze_valid"])].copy()
        h = pd.to_numeric(g.get("horizontal_visual_angle_deg"), errors="coerce")
        v = pd.to_numeric(g.get("vertical_visual_angle_deg"), errors="coerce")
        m = h.notna() & v.notna()
        h, v = h[m], v[m]
        fig, ax = plt.subplots(figsize=(7.2, 5.4))
        if len(h):
            hb = ax.hexbin(h, v, gridsize=35, mincnt=1, extent=full_extent)
            fig.colorbar(hb, ax=ax, label="Valid gaze samples per hexbin")
        ax.axhline(0, linewidth=0.8, linestyle="--")
        ax.axvline(0, linewidth=0.8, linestyle="--")
        ax.set_xlim(-hl,hl)
        ax.set_ylim(-vl,vl)
        ax.set_xlabel("Horizontal visual angle (deg; right +)")
        ax.set_ylabel("Vertical visual angle (deg; up +)")
        ax.set_title(f"{st['stage_id']} | {st['stage_type']}")
        ax.grid(True, alpha=0.25)
        fig.tight_layout()
        fig.savefig(figdir / f"stage_hv__{_safe_name(st['stage_id'])}__{_safe_name(st['stage_type'])}.png", dpi=160)
        plt.close(fig)

    # Primary AOI access timeline. Solid=episode containing exact hit; dashed=tolerance-only.
    if not primary_episodes.empty:
        p = primary_episodes.sort_values("start_s")
        order = list(dict.fromkeys(p["semantic_key"].astype(str).tolist()))
        ypos = {sem: i for i, sem in enumerate(order)}
        fig_h = max(4.0, 0.45 * len(order) + 2.5)
        fig, ax = plt.subplots(figsize=(10.5, fig_h))
        for _, r in p.iterrows():
            y = ypos[str(r["semantic_key"])]
            style = "-" if r["visit_hit_type"] == "exact" else "--"
            ax.plot([r["start_s"], r["end_s"]], [y, y], linewidth=4, linestyle=style)
            if r["duration_s"] <= 0.03:
                ax.plot([r["start_s"]], [y], marker="o", markersize=4)
        for _, st in stages.iterrows():
            ax.axvline(float(st["start_s"]), linewidth=0.8, linestyle=":")
            ax.text(float(st["start_s"]), len(order)-0.15, str(st["stage_id"]), rotation=90, va="top", ha="right", fontsize=8)
        if len(stages):
            ax.axvline(float(stages.iloc[-1]["end_s"]), linewidth=0.8, linestyle=":")
        ax.set_yticks(range(len(order)))
        ax.set_yticklabels(order)
        ax.set_xlabel("Video time (s)")
        ax.set_ylabel("Primary AOI semantic key")
        ax.set_title("AOI access timeline (solid = exact-containing visit; dashed = tolerance-only visit)")
        ax.grid(True, axis="x", alpha=0.25)
        fig.tight_layout()
        fig.savefig(figdir / "aoi_access_timeline.png", dpi=160)
        plt.close(fig)


def build_result_layer(
    fused: pd.DataFrame,
    long_hits: pd.DataFrame,
    stages: pd.DataFrame,
    stage_qc: pd.DataFrame,
    tracks: pd.DataFrame,
    outdir: str | Path,
    *,
    episode_gap_s: float = 0.10,
    scene_hfov_deg: float = 95.0,
    scene_vfov_deg: float = 63.0,
) -> dict:
    """Create analysis-ready summaries and figures from the validated fused layer."""
    outdir = Path(outdir)
    sample_period = _sample_period_s(fused)

    dedup_hits = _deduplicate_semantic_hits(long_hits, fused)
    # carry dataset IDs from fused because long hit rows intentionally stay compact in V0.2
    if not dedup_hits.empty:
        if "participant_id" not in dedup_hits.columns:
            dedup_hits.insert(0, "participant_id", fused["participant_id"].dropna().iloc[0] if fused["participant_id"].notna().any() else pd.NA)
        if "task_id" not in dedup_hits.columns:
            dedup_hits.insert(1, "task_id", fused["task_id"].dropna().iloc[0] if fused["task_id"].notna().any() else pd.NA)

    episodes = _build_semantic_episodes(dedup_hits, episode_gap_s, sample_period)
    primary_episodes = _build_primary_episodes(fused, episode_gap_s, sample_period)
    qc = _qc_summary(fused, stages, stage_qc)
    stage_summary = _stage_summary(fused, stages)
    aoi_summary = _aoi_summary(episodes, tracks)
    aoi_by_stage = _aoi_by_stage(episodes)
    transitions = _transitions(primary_episodes)
    doorplates = _doorplate_summary(episodes, tracks)

    # machine-readable result CSVs
    episodes.to_csv(outdir / "aoi_visits.csv", index=False, encoding="utf-8-sig")
    primary_episodes.to_csv(outdir / "aoi_primary_sequence.csv", index=False, encoding="utf-8-sig")
    qc.to_csv(outdir / "qc_summary.csv", index=False, encoding="utf-8-sig")
    stage_summary.to_csv(outdir / "stage_summary.csv", index=False, encoding="utf-8-sig")
    aoi_summary.to_csv(outdir / "aoi_summary.csv", index=False, encoding="utf-8-sig")
    aoi_by_stage.to_csv(outdir / "aoi_by_stage.csv", index=False, encoding="utf-8-sig")
    transitions.to_csv(outdir / "aoi_transition.csv", index=False, encoding="utf-8-sig")
    doorplates.to_csv(outdir / "doorplate_summary.csv", index=False, encoding="utf-8-sig")

    xlsx_path = outdir / "analysis_results.xlsx"
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        qc.to_excel(writer, sheet_name="QC_summary", index=False)
        stage_summary.to_excel(writer, sheet_name="Stage_summary", index=False)
        aoi_summary.to_excel(writer, sheet_name="AOI_summary", index=False)
        aoi_by_stage.to_excel(writer, sheet_name="AOI_by_stage", index=False)
        episodes.to_excel(writer, sheet_name="AOI_visits", index=False)
        transitions.to_excel(writer, sheet_name="AOI_transition", index=False)
        doorplates.to_excel(writer, sheet_name="Doorplate_summary", index=False)
    _style_excel(xlsx_path)

    _make_figures(fused, stages, primary_episodes, outdir, scene_hfov_deg, scene_vfov_deg)

    note = (
        "Result-layer definitions\n"
        "========================\n"
        f"AOI visit episode gap threshold: {episode_gap_s:.3f} s.\n"
        "AOI visit is an access episode, not a fixation.\n"
        "exact = gaze sample inside original CVAT rectangle.\n"
        "V1.1 exact only; tolerance compatibility columns contain no tolerance hits.\n"
        "Doorplate main semantics merge physical protruding/centered tracks by target_status + room_id.\n"
        "AOI_transition uses the non-overlapping primary AOI sequence; parent/child simultaneous hits are not double-counted there.\n"
        "Stage H/V figures use valid gaze samples; use 03 group baseline for common coordinate comparisons.\n"
    )
    (outdir / "RESULTS_README.txt").write_text(note, encoding="utf-8")

    return {
        "episode_gap_s": episode_gap_s,
        "sample_period_s_estimate": sample_period,
        "aoi_visit_count": int(len(episodes)),
        "primary_sequence_visit_count": int(len(primary_episodes)),
        "aoi_semantic_count_visited": int(aoi_summary["visited"].fillna(False).sum()) if len(aoi_summary) else 0,
        "doorplate_semantic_count_visited": int(doorplates["visited"].fillna(False).sum()) if len(doorplates) else 0,
        "analysis_results_xlsx": str(xlsx_path),
    }



def _main_scan_outputs():
    import sys
    root = Path(__file__).resolve().parent
    output_root = root / "output"
    config_path = root / "config" / "analysis_settings.json"
    if not output_root.exists():
        print("[停止] 没有找到 output/ 文件夹。请先运行 01_run_analysis.py。")
        raise SystemExit(1)

    try:
        import json
        settings = json.loads(config_path.read_text(encoding="utf-8")) if config_path.exists() else {}
    except Exception:
        settings = {}

    episode_gap_s = float(settings.get("aoi_episode_gap_s", 0.10))
    scene_hfov_deg = float(settings.get("scene_hfov_deg", 95.0))
    scene_vfov_deg = float(settings.get("scene_vfov_deg", 63.0))

    required = [
        "gaze_stage_aoi_fused.csv", "aoi_hits_effective_long.csv",
        "stages_standardized.csv", "stage_qc.csv", "aoi_tracks_standardized.csv"
    ]
    datasets = []
    for folder in sorted(p for p in output_root.iterdir() if p.is_dir() and not p.name.startswith(".")):
        missing = [fn for fn in required if not (folder / fn).exists()]
        if not missing:
            datasets.append(folder)

    if not datasets:
        print("[停止] output/ 里没有可生成结果层的数据包。请先运行 01_run_analysis.py。")
        raise SystemExit(1)

    print("\n=== G-line V0.3 结果层 ===")
    print(f"找到 {len(datasets)} 个已完成底层融合的数据包。")
    print(f"AOI visit gap = {episode_gap_s:.3f}s；exact / tolerance 分开保留。\n")

    rows = []
    for i, folder in enumerate(datasets, 1):
        print(f"[{i}/{len(datasets)}] {folder.name}")
        try:
            fused = pd.read_csv(folder / "gaze_stage_aoi_fused.csv")
            hits = pd.read_csv(folder / "aoi_hits_effective_long.csv")
            stages = pd.read_csv(folder / "stages_standardized.csv")
            sqc = pd.read_csv(folder / "stage_qc.csv")
            tracks = pd.read_csv(folder / "aoi_tracks_standardized.csv")
            result = build_result_layer(
                fused, hits, stages, sqc, tracks, folder,
                episode_gap_s=episode_gap_s,
                scene_hfov_deg=scene_hfov_deg,
                scene_vfov_deg=scene_vfov_deg,
            )
            rows.append({"dataset": folder.name, "status": "OK", **result})
            print(f"  [OK] AOI visits={result['aoi_visit_count']}, visited semantics={result['aoi_semantic_count_visited']}")
        except ModuleNotFoundError as e:
            if e.name == "matplotlib":
                print("  [ERROR] 缺少 matplotlib。请在当前 (.venv) 终端运行: pip install matplotlib")
            else:
                print(f"  [ERROR] 缺少依赖: {e.name}")
            rows.append({"dataset": folder.name, "status": "ERROR", "message": str(e)})
        except Exception as e:
            import traceback
            print("  [ERROR]", e)
            (folder / "RESULT_LAYER_ERROR.txt").write_text(traceback.format_exc(), encoding="utf-8")
            rows.append({"dataset": folder.name, "status": "ERROR", "message": str(e)})

    pd.DataFrame(rows).to_csv(output_root / "result_layer_summary.csv", index=False, encoding="utf-8-sig")
    ok = sum(r.get("status") == "OK" for r in rows)
    err = sum(r.get("status") == "ERROR" for r in rows)
    print("\n=== 结果层完成 ===")
    print(f"成功: {ok}  失败: {err}")
    print("每个 output/被试__任务/ 下新增 analysis_results.xlsx、结果 CSV 和 figures/。")


if __name__ == "__main__":
    _main_scan_outputs()
